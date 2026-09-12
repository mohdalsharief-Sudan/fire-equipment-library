# app/projects_loader.py
"""
تحميل المشاريع من ملف Excel إلى قاعدة البيانات
"""

import os
import logging
from openpyxl import load_workbook

from .database import EquipmentDatabase

logger = logging.getLogger(__name__)


class ProjectsLoader:
    """تحميل المشاريع من Excel"""
    
    def __init__(self, db: EquipmentDatabase):
        self.db = db
    
    def load_from_excel(self, file_path: str) -> dict:
        """تحميل المشاريع من ملف Excel"""
        
        if not os.path.exists(file_path):
            logger.error(f"❌ الملف غير موجود: {file_path}")
            return {'projects': 0, 'clients': 0}
        
        wb = load_workbook(file_path, data_only=True)
        ws = wb['ملخص المشروعات']
        
        projects_count = 0
        clients_count = 0
        
        # طباعة الأعمدة للتشخيص
        print("🔍 تشخيص الأعمدة (الصف 5):")
        for i, cell in enumerate(ws[5]):
            print(f"   عمود {i}: {cell.value}")
        print()
        
        # الصفوف 6-23 تحتوي على المشاريع
        for row_num in range(6, 24):
            row = ws[row_num]
            cells = list(row)
            
            # قراءة الأعمدة
            
            project_name = cells[2].value if len(cells) > 2 else None  # اسم المشروع
            client_name = cells[3].value if len(cells) > 3 else None   # العميل
            fire_cost = cells[4].value if len(cells) > 4 else 0        # إطفاء
            alarm_cost = cells[5].value if len(cells) > 5 else 0       # إنذار
            vent_cost = cells[6].value if len(cells) > 6 else 0        # تهوية
            other_cost = cells[7].value if len(cells) > 7 else 0       # أخرى
            
            if not project_name:
                continue
            
            # تحويل إلى float
            def to_float(v):
                if v is None:
                    return 0
                try:
                    return float(v)
                except (ValueError, TypeError):
                    return 0
            
            fire_cost = to_float(fire_cost)
            alarm_cost = to_float(alarm_cost)
            vent_cost = to_float(vent_cost)
            other_cost = to_float(other_cost)
            
            # إضافة العميل
            client_id = None
            if client_name:
                client_id = self.db.add_client(
                    name=str(client_name).strip()[:200],
                    notes="مستورد من ملف Excel"
                )
                clients_count += 1
            
            # حساب الإجماليات
            total_before_vat = fire_cost + alarm_cost + vent_cost + other_cost
            vat_15 = total_before_vat * 0.15
            total_with_vat = total_before_vat + vat_15
            
            # إضافة المشروع
            self.db.add_project({
                'project_name': str(project_name).strip()[:200],
                'client_id': client_id,
                'service_provider': 'مؤسسة أسس السلامة / رئال الدولية',
                'fire_suppression_cost': fire_cost,
                'fire_alarm_cost': alarm_cost,
                'ventilation_cost': vent_cost,
                'other_systems_cost': other_cost,
                'total_before_vat': total_before_vat,
                'vat_15': vat_15,
                'total_with_vat': total_with_vat,
                'notes': 'مستورد من Excel',
                'status': 'completed'
            })
            
            projects_count += 1
        
        logger.info(f"✅ تم تحميل {projects_count} مشروع و {clients_count} عميل")
        
        return {
            'projects': projects_count,
            'clients': clients_count
        }