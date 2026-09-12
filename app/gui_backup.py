# app/gui.py
"""
واجهة رسومية لمكتبة معدات السلامة
"""

import sys
import os
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
import threading

sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))

from app.search import EquipmentSearch


class EquipmentLibraryGUI:
    """الواجهة الرسومية لمكتبة المعدات"""
    
    def __init__(self):
        self.root = tk.Tk()
        self.root.title("Fire Equipment Library - مكتبة معدات السلامة")
        self.root.geometry("1200x750")
        self.root.configure(bg="#0f1626")
        
        # المتغيرات
        self.search_var = tk.StringVar()
        self.manufacturer_var = tk.StringVar(value="الكل")
        self.category_var = tk.StringVar(value="الكل")
        self.min_flow_var = tk.StringVar()
        self.max_price_var = tk.StringVar()
        
        # البيانات
        self.search = EquipmentSearch()
        self.results = []
        
        self._build_ui()
        self._load_all()
    
    def _build_ui(self):
        """بناء الواجهة"""
        # ====== العنوان ======
        title_frame = tk.Frame(self.root, bg="#1B4F72")
        title_frame.pack(fill="x", padx=10, pady=10)
        
        title_label = tk.Label(
            title_frame,
            text="🔥 Fire Equipment Library",
            font=("Arial", 22, "bold"),
            bg="#1B4F72",
            fg="white",
            padx=20,
            pady=12,
        )
        title_label.pack()
        
        subtitle = tk.Label(
            title_frame,
            text="مكتبة معدات السلامة - SFFECO | NAFFCO",
            font=("Arial", 11),
            bg="#1B4F72",
            fg="#A8D5E5",
        )
        subtitle.pack(pady=(0, 8))
        
        # ====== إطار البحث ======
        search_frame = tk.Frame(self.root, bg="#0f1626")
        search_frame.pack(fill="x", padx=10, pady=5)
        
        tk.Label(
            search_frame, text="🔍 بحث:",
            font=("Arial", 11, "bold"),
            bg="#0f1626", fg="white"
        ).pack(side="left", padx=5)
        
        search_entry = tk.Entry(
            search_frame, textvariable=self.search_var,
            width=30, font=("Arial", 11)
        )
        search_entry.pack(side="left", padx=5)
        search_entry.bind("<Return>", lambda e: self._do_search())
        
        # المصنع
        tk.Label(
            search_frame, text="المصنع:",
            font=("Arial", 10), bg="#0f1626", fg="white"
        ).pack(side="left", padx=5)
        
        manufacturer_combo = ttk.Combobox(
            search_frame, textvariable=self.manufacturer_var,
            values=["الكل", "SFFECO", "NAFFCO"],
            width=12, font=("Arial", 10), state="readonly"
        )
        manufacturer_combo.pack(side="left", padx=5)
        
        # الفئة
        tk.Label(
            search_frame, text="الفئة:",
            font=("Arial", 10), bg="#0f1626", fg="white"
        ).pack(side="left", padx=5)
        
        category_combo = ttk.Combobox(
            search_frame, textvariable=self.category_var,
            values=["الكل", "package_units", "single_pumps", "split_case", "jockey", "controllers", "valves", "tanks"],
            width=15, font=("Arial", 10), state="readonly"
        )
        category_combo.pack(side="left", padx=5)
        
        # زر البحث
        search_btn = tk.Button(
            search_frame, text="🔍 بحث",
            command=self._do_search,
            bg="#27AE60", fg="white",
            font=("Arial", 10, "bold"),
            padx=15, pady=3, cursor="hand2"
        )
        search_btn.pack(side="left", padx=10)
        
        # زر عرض الكل
        all_btn = tk.Button(
            search_frame, text="📋 عرض الكل",
            command=self._load_all,
            bg="#2874A6", fg="white",
            font=("Arial", 10, "bold"),
            padx=15, pady=3, cursor="hand2"
        )
        all_btn.pack(side="left", padx=5)
        
        # ====== فلاتر متقدمة ======
        filter_frame = tk.Frame(self.root, bg="#0f1626")
        filter_frame.pack(fill="x", padx=10, pady=5)
        
        tk.Label(
            filter_frame, text="أدنى تدفق (GPM):",
            font=("Arial", 10), bg="#0f1626", fg="#A8D5E5"
        ).pack(side="left", padx=5)
        
        tk.Entry(
            filter_frame, textvariable=self.min_flow_var,
            width=10, font=("Arial", 10)
        ).pack(side="left", padx=5)
        
        tk.Label(
            filter_frame, text="أقصى سعر (ريال):",
            font=("Arial", 10), bg="#0f1626", fg="#A8D5E5"
        ).pack(side="left", padx=5)
        
        tk.Entry(
            filter_frame, textvariable=self.max_price_var,
            width=12, font=("Arial", 10)
        ).pack(side="left", padx=5)
        
        # ====== جدول النتائج ======
        table_frame = tk.Frame(self.root, bg="#0f1626")
        table_frame.pack(fill="both", expand=True, padx=10, pady=10)
        
        # Scrollbars
        scroll_y = tk.Scrollbar(table_frame, orient="vertical")
        scroll_x = tk.Scrollbar(table_frame, orient="horizontal")
        
        # Treeview
        columns = ("#", "manufacturer", "model", "type", "flow", "pressure", "price", "package")
        self.tree = ttk.Treeview(
            table_frame, columns=columns, show="headings",
            yscrollcommand=scroll_y.set, xscrollcommand=scroll_x.set
        )
        
        # العناوين
        headers = {
            "#": "#",
            "manufacturer": "المصنع",
            "model": "الموديل",
            "type": "النوع",
            "flow": "التدفق (GPM)",
            "pressure": "الضغط (bar)",
            "price": "السعر (ريال)",
            "package": "الحزمة"
        }
        
        widths = {
            "#": 40,
            "manufacturer": 100,
            "model": 180,
            "type": 180,
            "flow": 100,
            "pressure": 100,
            "price": 120,
            "package": 150
        }
        
        for col in columns:
            self.tree.heading(col, text=headers[col])
            self.tree.column(col, width=widths[col], anchor="center")
        
        self.tree.pack(side="left", fill="both", expand=True)
        scroll_y.config(command=self.tree.yview)
        scroll_x.config(command=self.tree.xview)
        scroll_y.pack(side="right", fill="y")
        scroll_x.pack(side="bottom", fill="x")
        
        # Bind للنقر
        self.tree.bind("<Double-1>", self._on_item_double_click)
        
        # ====== شريط الحالة ======
        status_frame = tk.Frame(self.root, bg="#1B4F72")
        status_frame.pack(fill="x", side="bottom")
        
        self.status_var = tk.StringVar(value="جاهز")
        tk.Label(
            status_frame, textvariable=self.status_var,
            font=("Arial", 10), bg="#1B4F72", fg="white",
            padx=10, pady=5
        ).pack(side="left")
        
        self.count_var = tk.StringVar(value="0 معدة")
        tk.Label(
            status_frame, textvariable=self.count_var,
            font=("Arial", 10, "bold"), bg="#1B4F72", fg="#A8D5E5",
            padx=10, pady=5
        ).pack(side="right")
        
        # ====== أزرار التصدير ======
        export_frame = tk.Frame(self.root, bg="#0f1626")
        export_frame.pack(fill="x", padx=10, pady=5)
        
        tk.Button(
            export_frame, text="📊 تصدير Excel",
            command=self._export_excel,
            bg="#27AE60", fg="white",
            font=("Arial", 10, "bold"),
            padx=15, pady=5, cursor="hand2"
        ).pack(side="left", padx=5)
        
        tk.Button(
            export_frame, text="📄 تصدير JSON",
            command=self._export_json,
            bg="#8E44AD", fg="white",
            font=("Arial", 10, "bold"),
            padx=15, pady=5, cursor="hand2"
        ).pack(side="left", padx=5)
    
    def _load_all(self):
        """عرض كل المعدات"""
        self.status_var.set("جاري التحميل...")
        self.results = self.search.list_all()
        self._display_results()
    
    def _do_search(self):
        """تنفيذ البحث"""
        self.status_var.set("جاري البحث...")
        
        search_text = self.search_var.get().strip()
        manufacturer = self.manufacturer_var.get()
        category = self.category_var.get()
        
        # استعلام أساسي
        min_flow = None
        if self.min_flow_var.get().strip():
            try:
                min_flow = float(self.min_flow_var.get())
            except ValueError:
                pass
        
        max_price = None
        if self.max_price_var.get().strip():
            try:
                max_price = float(self.max_price_var.get())
            except ValueError:
                pass
        
        # استخدام find_pumps إذا البحث عن مضخات
        if category in ["الكل", "package_units", "single_pumps", "split_case", "jockey"]:
            self.results = self.search.find_pumps(
                flow_gpm=min_flow,
                manufacturer=manufacturer if manufacturer != "الكل" else None,
                max_price=max_price
            )
        else:
            # استعلام عام
            self.results = self.search.list_all()
            if manufacturer != "الكل":
                self.results = [r for r in self.results if r['manufacturer_name'] == manufacturer]
            if max_price:
                self.results = [r for r in self.results if r.get('price_sar', 0) <= max_price]
        
        # فلترة حسب نص البحث
        if search_text:
            search_lower = search_text.lower()
            self.results = [
                r for r in self.results
                if search_lower in r.get('model', '').lower()
                or search_lower in r.get('manufacturer_name', '').lower()
                or search_lower in r.get('type', '').lower()
            ]
        
        self._display_results()
    
    def _display_results(self):
        """عرض النتائج في الجدول"""
        # مسح القديم
        for item in self.tree.get_children():
            self.tree.delete(item)
        
        # إضافة النتائج
        for i, eq in enumerate(self.results, 1):
            specs = eq.get('specs', {})
            
            # العمود: الحزمة
            package_info = ""
            if specs.get('is_complete_package'):
                package_info = "✅ كاملة"
            elif specs.get('includes_jockey'):
                package_info = "✅ مع جوكي"
            elif specs.get('requires_jockey'):
                package_info = "⚠️ تحتاج جوكي"
            
            self.tree.insert("", "end", values=(
                i,
                eq.get('manufacturer_name', ''),
                eq.get('model', ''),
                eq.get('type', ''),
                specs.get('flow_gpm', 0),
                specs.get('pressure_bar', 0),
                f"{eq.get('price_sar', 0):,.0f}",
                package_info
            ))
        
        self.count_var.set(f"{len(self.results)} معدة")
        self.status_var.set("جاهز")
    
    def _on_item_double_click(self, event):
        """عرض تفاصيل المعدة عند النقر المزدوج"""
        selection = self.tree.selection()
        if not selection:
            return
        
        item = self.tree.item(selection[0])
        index = int(item['values'][0]) - 1
        
        if 0 <= index < len(self.results):
            equipment = self.results[index]
            self._show_details(equipment)
    
    def _show_details(self, equipment):
        """عرض تفاصيل معدة في نافذة منبثقة"""
        window = tk.Toplevel(self.root)
        window.title(f"تفاصيل: {equipment['model']}")
        window.geometry("600x600")
        window.configure(bg="#0f1626")
        
        # العنوان
        tk.Label(
            window,
            text=f"🔧 {equipment['manufacturer_name']} - {equipment['model']}",
            font=("Arial", 16, "bold"),
            bg="#0f1626", fg="#3498DB",
            pady=15
        ).pack()
        
        # المحتوى
        text_frame = tk.Frame(window, bg="#0f1626")
        text_frame.pack(fill="both", expand=True, padx=20, pady=10)
        
        text_widget = tk.Text(
            text_frame,
            font=("Arial", 11),
            bg="#1a1a2e", fg="#e0e0e0",
            wrap="word", padx=15, pady=15
        )
        text_widget.pack(fill="both", expand=True)
        button_frame = tk.Frame(window, bg="#0f1626")
        button_frame.pack(pady=10)
        
        tk.Button(
            button_frame,
            text="💰 تعديل السعر",
            command=lambda: self._edit_price(equipment, window),
            bg="#E67E22", fg="white",
            font=("Arial", 11, "bold"),
            padx=20, pady=8, cursor="hand2"
        ).pack(side="left", padx=5)
        
        tk.Button(
            button_frame,
            text="📊 تاريخ الأسعار",
            command=lambda: self._show_price_history(equipment),
            bg="#8E44AD", fg="white",
            font=("Arial", 11, "bold"),
            padx=20, pady=8, cursor="hand2"
        ).pack(side="left", padx=5)
        
        tk.Button(
            button_frame,
            text="❌ إغلاق",
            command=window.destroy,
            bg="#E74C3C", fg="white",
            font=("Arial", 11, "bold"),
            padx=20, pady=8, cursor="hand2"
        ).pack(side="left", padx=5)
        # بناء النص
        specs = equipment.get('specs', {})
        lines = []
        lines.append(f"النوع: {equipment.get('type', '')}\n")
        
        if specs.get('flow_gpm'):
            lines.append(f"التدفق: {specs['flow_gpm']} GPM\n")
        if specs.get('pressure_bar'):
            lines.append(f"الضغط: {specs['pressure_bar']} bar\n")
        if specs.get('power_hp'):
            lines.append(f"القدرة: {specs['power_hp']} HP\n")
        if specs.get('electric_hp'):
            lines.append(f"مضخة كهربائية: {specs['electric_hp']} HP\n")
        if specs.get('diesel_hp'):
            lines.append(f"مضخة ديزل: {specs['diesel_hp']} HP\n")
        if specs.get('jockey_hp'):
            lines.append(f"مضخة جوكي: {specs['jockey_hp']} HP\n")
        if specs.get('weight_kg'):
            lines.append(f"الوزن: {specs['weight_kg']} kg\n")
        if specs.get('header_pipe'):
            lines.append(f"الهيدر: {specs['header_pipe']}\n")
        
        lines.append(f"\n📦 محتويات الحزمة:\n")
        if specs.get('is_complete_package'):
            lines.append("   ✅ حزمة كاملة جاهزة\n")
        if specs.get('includes_jockey'):
            lines.append("   ✅ تشمل مضخة جوكي\n")
        if specs.get('includes_controller'):
            lines.append("   ✅ تشمل لوحة تحكم\n")
        if specs.get('includes_diesel'):
            lines.append("   ✅ تشمل مضخة ديزل\n")
        
        if equipment.get('certifications'):
            lines.append(f"\n🏆 الاعتمادات: {' / '.join(equipment['certifications'])}\n")
        
        lines.append(f"\n💰 السعر: {equipment.get('price_sar', 0):,.0f} ريال\n")
        
        if equipment.get('applications'):
            lines.append(f"\n📋 التطبيقات:\n")
            for app in equipment['applications']:
                lines.append(f"   • {app}\n")
        
        if equipment.get('notes'):
            lines.append(f"\n📝 ملاحظات: {equipment['notes']}\n")
        
        text_widget.insert("1.0", "".join(lines))
        text_widget.config(state="disabled")
        
        # زر إغلاق
        tk.Button(
            window, text="إغلاق",
            command=window.destroy,
            bg="#E74C3C", fg="white",
            font=("Arial", 11, "bold"),
            padx=30, pady=8, cursor="hand2"
        ).pack(pady=10)
    
    def _edit_price(self, equipment, parent_window):
        """تعديل سعر معدة"""
        # نافذة إدخال
        price_window = tk.Toplevel(parent_window)
        price_window.title("تعديل السعر")
        price_window.geometry("400x250")
        price_window.configure(bg="#0f1626")
        
        tk.Label(
            price_window,
            text=f"🔧 {equipment['model']}",
            font=("Arial", 14, "bold"),
            bg="#0f1626", fg="#3498DB",
            pady=15
        ).pack()
        
        tk.Label(
            price_window,
            text=f"السعر الحالي: {equipment.get('price_sar', 0):,.0f} ريال",
            font=("Arial", 11),
            bg="#0f1626", fg="white"
        ).pack(pady=5)
        
        tk.Label(
            price_window,
            text="السعر الجديد (ريال):",
            font=("Arial", 11, "bold"),
            bg="#0f1626", fg="#A8D5E5"
        ).pack(pady=10)
        
        price_var = tk.StringVar(value=str(equipment.get('price_sar', 0)))
        price_entry = tk.Entry(
            price_window,
            textvariable=price_var,
            font=("Arial", 14),
            width=15,
            justify="center"
        )
        price_entry.pack(pady=5)
        price_entry.select_range(0, "end")
        price_entry.focus()
        
        def save_price():
            try:
                new_price = float(price_var.get())
                if new_price < 0:
                    messagebox.showerror("خطأ", "السعر يجب أن يكون موجباً")
                    return
                
                # الحصول على ID المعدة من قاعدة البيانات
                equipment_id = self._get_equipment_id(equipment['model'])
                
                if equipment_id:
                    # تحديث في قاعدة البيانات
                    self.search.conn.execute("""
                        UPDATE equipment SET price_sar = ? WHERE id = ?
                    """, (new_price, equipment_id))
                    self.search.conn.commit()
                    
                    # تحديث في الذاكرة
                    equipment['price_sar'] = new_price
                    
                    # إعادة عرض النتائج
                    self._display_results()
                    
                    messagebox.showinfo("تم", f"✅ تم تحديث السعر إلى {new_price:,.0f} ريال")
                    price_window.destroy()
                    parent_window.destroy()
                    
                    # فتح التفاصيل مرة أخرى بالسعر الجديد
                    self._show_details(equipment)
                else:
                    messagebox.showerror("خطأ", "لم يتم العثور على المعدة")
                    
            except ValueError:
                messagebox.showerror("خطأ", "أدخل رقماً صحيحاً")
        
        tk.Button(
            price_window,
            text="💾 حفظ",
            command=save_price,
            bg="#27AE60", fg="white",
            font=("Arial", 11, "bold"),
            padx=30, pady=8, cursor="hand2"
        ).pack(pady=15)
        
        price_entry.bind("<Return>", lambda e: save_price())
    
    def _get_equipment_id(self, model: str) -> int:
        """الحصول على ID المعدة من قاعدة البيانات"""
        cursor = self.search.conn.cursor()
        cursor.execute("SELECT id FROM equipment WHERE model = ?", (model,))
        row = cursor.fetchone()
        return row[0] if row else None
    
    def _show_price_history(self, equipment):
        """عرض تاريخ الأسعار"""
        equipment_id = self._get_equipment_id(equipment['model'])
        if not equipment_id:
            messagebox.showerror("خطأ", "لم يتم العثور على المعدة")
            return
        
        cursor = self.search.conn.cursor()
        cursor.execute("""
            SELECT price_sar, source, changed_at
            FROM price_history
            WHERE equipment_id = ?
            ORDER BY changed_at DESC
        """, (equipment_id,))
        
        history = cursor.fetchall()
        
        # نافذة التاريخ
        hist_window = tk.Toplevel(self.root)
        hist_window.title(f"تاريخ الأسعار - {equipment['model']}")
        hist_window.geometry("500x400")
        hist_window.configure(bg="#0f1626")
        
        tk.Label(
            hist_window,
            text=f"📊 تاريخ أسعار {equipment['model']}",
            font=("Arial", 14, "bold"),
            bg="#0f1626", fg="#3498DB",
            pady=15
        ).pack()
        
        # قائمة
        listbox = tk.Listbox(
            hist_window,
            font=("Consolas", 11),
            bg="#1a1a2e", fg="#e0e0e0",
            width=60, height=15
        )
        listbox.pack(padx=20, pady=10, fill="both", expand=True)
        
        for row in history:
            price, source, date = row
            listbox.insert(tk.END, f"{date[:16]} | {price:>10,.0f} ريال | {source}")
        
        if not history:
            listbox.insert(tk.END, "لا يوجد سجل أسعار بعد")
        
        tk.Button(
            hist_window,
            text="إغلاق",
            command=hist_window.destroy,
            bg="#E74C3C", fg="white",
            font=("Arial", 11, "bold"),
            padx=30, pady=8, cursor="hand2"
        ).pack(pady=10)
    
    def _export_excel(self):
        """تصدير Excel"""
        if not self.results:
            messagebox.showwarning("تنبيه", "لا توجد نتائج للتصدير")
            return
        
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment
        except ImportError:
            messagebox.showerror("خطأ", "مكتبة openpyxl غير مثبتة")
            return
        
        # ← المسار الافتراضي: exports/
        exports_dir = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'exports')
        os.makedirs(exports_dir, exist_ok=True)
        
        from datetime import datetime
        default_name = f"equipment_export_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        
        filepath = filedialog.asksaveasfilename(
            title="حفظ Excel",
            initialdir=exports_dir,           # ← يبدأ من exports
            initialfile=default_name,          # ← اسم افتراضي
            defaultextension=".xlsx",
            filetypes=[("Excel", "*.xlsx")]
        )
        if not filepath:
            return
        
        
        wb = Workbook()
        ws = wb.active
        ws.title = "المعدات"
        ws.sheet_view.rightToLeft = True
        
        # العناوين
        headers = ["#", "المصنع", "الموديل", "النوع", "التدفق (GPM)", "الضغط (bar)", "السعر (ريال)", "الحزمة"]
        ws.append(headers)
        
        # تنسيق العناوين
        header_fill = PatternFill(start_color="1B4F72", end_color="1B4F72", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=12)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        
        # البيانات
        for i, eq in enumerate(self.results, 1):
            specs = eq.get('specs', {})
            package = ""
            if specs.get('is_complete_package'):
                package = "كاملة"
            elif specs.get('includes_jockey'):
                package = "مع جوكي"
            
            ws.append([
                i,
                eq.get('manufacturer_name', ''),
                eq.get('model', ''),
                eq.get('type', ''),
                specs.get('flow_gpm', 0),
                specs.get('pressure_bar', 0),
                eq.get('price_sar', 0),
                package
            ])
        
        # ضبط عرض الأعمدة
        widths = [5, 15, 25, 25, 15, 15, 18, 15]
        for i, width in enumerate(widths, 1):
            ws.column_dimensions[chr(64 + i)].width = width
        
        wb.save(filepath)
        messagebox.showinfo("تم", f"تم الحفظ:\n{filepath}")
    
    def _export_json(self):
        """تصدير JSON"""
        if not self.results:
            messagebox.showwarning("تنبيه", "لا توجد نتائج للتصدير")
            return
        
        import json
        from datetime import datetime
        
        # ← المسار الافتراضي: exports/
        exports_dir = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'exports')
        os.makedirs(exports_dir, exist_ok=True)
        
        default_name = f"equipment_export_{datetime.now().strftime('%Y%m%d_%H%M%S')}.json"
        
        filepath = filedialog.asksaveasfilename(
            title="حفظ JSON",
            initialdir=exports_dir,
            initialfile=default_name,
            defaultextension=".json",
            filetypes=[("JSON", "*.json")]
        )
        if not filepath:
            return
        
        
        # تبسيط البيانات
        export_data = []
        for eq in self.results:
            export_data.append({
                'manufacturer': eq.get('manufacturer_name', ''),
                'model': eq.get('model', ''),
                'type': eq.get('type', ''),
                'specs': eq.get('specs', {}),
                'price_sar': eq.get('price_sar', 0),
                'certifications': eq.get('certifications', []),
            })
        
        with open(filepath, 'w', encoding='utf-8') as f:
            json.dump(export_data, f, ensure_ascii=False, indent=2)
        
        messagebox.showinfo("تم", f"تم الحفظ:\n{filepath}")
    
    def run(self):
        """تشغيل الواجهة"""
        self.root.mainloop()


def main():
    """الدالة الرئيسية"""
    gui = EquipmentLibraryGUI()
    gui.run()


if __name__ == "__main__":
    main()