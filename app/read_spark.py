# app/read_spark.py
"""
قراءة ملف مشروع SPARK
"""

import os
from openpyxl import load_workbook


def read_excel_file(file_path):
    """قراءة ملف Excel بشكل تفصيلي"""
    
    if not os.path.exists(file_path):
        print(f"[X] الملف غير موجود: {file_path}")
        return
    
    print(f"[*] قراءة: {file_path}")
    print("=" * 80)
    
    wb = load_workbook(file_path, data_only=True)
    
    print(f"[*] عدد الأوراق: {len(wb.worksheets)}")
    print()
    
    for sheet_idx, sheet_name in enumerate(wb.sheetnames, 1):
        ws = wb[sheet_name]
        print(f"\n{'=' * 80}")
        print(f"[ورقة {sheet_idx}] {sheet_name}")
        print(f"   الأبعاد: {ws.max_row} صف × {ws.max_column} عمود")
        print('=' * 80)
        
        # عرض أول 40 صف
        for i, row in enumerate(ws.iter_rows(max_row=40, values_only=True), 1):
            if all(cell is None for cell in row):
                continue
            
            cells_display = []
            for cell in row:
                if cell is None:
                    cells_display.append("")
                elif isinstance(cell, (int, float)):
                    if isinstance(cell, float):
                        cells_display.append(f"{cell:,.2f}")
                    else:
                        cells_display.append(f"{cell:,}")
                else:
                    cells_display.append(str(cell)[:35])
            
            # إزالة الأعمدة الفارغة في النهاية
            while cells_display and cells_display[-1] == "":
                cells_display.pop()
            
            if cells_display:
                print(f"  صف {i:3}: {' | '.join(cells_display)}")
        
        if ws.max_row > 40:
            print(f"  ... ({ws.max_row - 40} صف إضافي)")
        print()


if __name__ == "__main__":
    # البحث عن ملف سبارك
    docs_dir = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'documents')
    
    # البحث في كل المجلدات الفرعية
    found = False
    for root, dirs, files in os.walk(docs_dir):
        for file in files:
            if 'سبارك' in file or 'spark' in file.lower():
                file_path = os.path.join(root, file)
                print(f"\n[✓] تم إيجاد الملف: {file_path}\n")
                read_excel_file(file_path)
                found = True
                break
        if found:
            break
    
    if not found:
        print("[X] لم يتم العثور على ملف سبارك")
        print(f"[*] البحث في: {docs_dir}")
        print("\n[*] الملفات الموجودة:")
        for root, dirs, files in os.walk(docs_dir):
            for file in files:
                print(f"   - {os.path.join(root, file)}")