# app/read_new_file.py
"""
قراءة ملف جدول البنود والمعدات المستقلة
"""

import os
from openpyxl import load_workbook

def read_excel_file(file_path):
    """قراءة ملف Excel وعرض محتواه"""
    
    if not os.path.exists(file_path):
        print(f"❌ الملف غير موجود: {file_path}")
        return
    
    print(f"📂 قراءة: {file_path}")
    print("=" * 70)
    
    wb = load_workbook(file_path, data_only=True)
    
    print(f"📊 عدد الأوراق: {len(wb.worksheets)}")
    print()
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        print(f"\n{'=' * 70}")
        print(f"📄 ورقة: {sheet_name}")
        print(f"   الأبعاد: {ws.max_row} صف × {ws.max_column} عمود")
        print('=' * 70)
        
        # عرض أول 50 صف
        for i, row in enumerate(ws.iter_rows(max_row=50, values_only=True), 1):
            if all(cell is None for cell in row):
                continue
            
            cells_display = []
            for cell in row:
                if cell is None:
                    cells_display.append("")
                elif isinstance(cell, (int, float)):
                    if isinstance(cell, float):
                        cells_display.append(f"{cell:,.2f}")
                    else:
                        cells_display.append(f"{cell:,}")
                else:
                    cells_display.append(str(cell)[:35])
            
            print(f"  صف {i:3}: {' | '.join(cells_display)}")
        
        if ws.max_row > 50:
            print(f"  ... ({ws.max_row - 50} صف إضافي)")


if __name__ == "__main__":
    file_path = r"D:\My-GitHub\GitHub\fire-equipment-library\data\projects\جدول_البنود_والمعدات_المستقلة_لأنظمة_السلامة.xlsx"
    read_excel_file(file_path)